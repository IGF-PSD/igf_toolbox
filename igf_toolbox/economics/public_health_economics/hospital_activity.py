import holidays
import numpy as np
import pandas as pd
from openpyxl import load_workbook


class HospitalActivityDiamant:
    """ 
    A class for computing hospital activity from DIAMANT data and INSEE regional demographic estimations.
    DIAMANT data should be in a proper format as given in the documentation.
    INSEE data should be the regional population estimations with age classes of 5 years.
    """

    def __init__(self, file_path_activity: str, file_path_demography: str):
        """ 
        Initializes the HospitalActivityDiamant object by loading various Excel sheets
        containing activity data and valuation data.

        Args:
            file_path_activity (str): Path to Excel file with activity and valuation DIAMANT data.
            file_path_demography (str): Path to Excel file with demographic data from INSEE.
        """

        self.data_valorisations = pd.read_excel(
            file_path_activity, sheet_name="Valorisations", header=4
        )
        self.data_valorisations_racine = pd.read_excel(
            file_path_activity, sheet_name="Valorisations - racine", header=4
        )
        self.data_valorisations_severite = pd.read_excel(
            file_path_activity, sheet_name="Valorisations - sévérité", header=5
        )
        self.data_valorisations_type_hosp = pd.read_excel(
            file_path_activity, sheet_name="Valorisations - type hosp", header=4
        )
        self.data_casemix = pd.read_excel(
            file_path_activity, sheet_name="Volume économique", header=1
        )
        self.file_path_demography = file_path_demography

        self.ghm = "PMSI MCO - GHM"
        self.ghs = "PMSI MCO - GHS"
        self.sejours_prix = "ACTIVITE - Nb Séjours Hors Séances"
        self.montants_am = "Montant AM GHS"
        self.taux_am = "Taux de remboursement AM"
        self.type_hosp = "PMSI MCO - Activité - Type Hospitalisation"
        self.tranche_age = "PMSI - Tranche Age INSEE"
        self.hosp_hp = "Ambulatoire"
        self.hosp_hc = "HC"
        self.racine = "PMSI MCO - GHM - Code Racine"
        self.severite = "PMSI MCO - GHM - Niveau de sévérité"
        self.list_severite = ["1", "2", "3", "4", "A", "B", "C", "D"]

    def _count_number_working_days(self, year: int) -> tuple[int]:
        """
        Returns the number of working and non working days in a given year.

        Args:
            year (int): The year for which numbers of days should be returned.

        Returns:
            tuple[int]: A tuple with a first element being number of working days and
                        second element being number of non working days.
        """

        # Define a list of dates of French holidays
        french_holidays_dates = holidays.France(years=year)

        # Count the number of working days and non working days (i.e weekends and holidays)
        dates = pd.date_range(start=f"{year}-01-01", end=f"{year}-12-31", freq="D")
        number_working_days = sum(
            1
            for date in dates
            if (date.weekday() < 5) and (date not in french_holidays_dates)
        )
        number_non_working_days = len(dates) - number_working_days

        return number_working_days, number_non_working_days

    def _effet_cjo(self, year: int, activity: str = "volume") -> float:
        """
        Returns the holiday adjustments to apply to yearly evolutions
        between year and year - 1.

        Args:
            year (int): The reference year for computing the effect in regards to the previous year.
            activity (str, optional): Type of weight to apply to the adjustment. 49% for `volume`
                                      and 34% for `sejours`. Values in [`volume`,`sejours`].
                                      Defaults to `volume`.

        Returns:
            float: CJO adjustment between year and year - 1.

        Raises:
            ValueError: Raises an error if `activity` parameter is not in [`volume`,`sejours`].
        """

        # Define the weight of the non working days
        if activity == "volume":
            coeff = 0.49
        elif activity == "sejours":
            coeff = 0.34
        else:
            raise ValueError("""
                            `activity` parameter should be in [`volume`, `séjours`] for holiday adjustments.
                            For `volume` a 49% correction is applied, for `séjours` a 34% correction is applied.
                            `séjours` is applied for both stays and `équivalents-journées`.
                            """)

        # Compute the number of `jours d'activite` and return their evolution
        # as CJO adjustment
        jours_activite_n = (
            self._count_number_working_days(year)[0]
            + coeff * self._count_number_working_days(year)[1]
        )
        jours_activite_n_1 = (
            self._count_number_working_days(year - 1)[0]
            + coeff * self._count_number_working_days(year - 1)[1]
        )

        return jours_activite_n / jours_activite_n_1 - 1

    @staticmethod
    def _preprocess_ghm(x: str) -> str:
        """ 
        Extracts the 6-character GHM code from a full string with format 'GHM - Label'.

        Args:
            x (str): Full GHM string in format 'GHM - Label'.

        Returns:
            str: 6-character GHM code.
        """
        x = x.split(" - ")[0]
        return x[0:6]

    @staticmethod
    def _preprocess_racine(x: str) -> str:
        """ 
        Extracts the 5-character GHM root code from a full string with format 'GHM - Label'.

        Args:
            x (str): Full GHM string in format 'GHM - Label'.

        Returns:
            str: 5-character root code.
        """
        x = x.split(" - ")[0]
        return x[0:5]

    @staticmethod
    def _preprocess_severite(x: str) -> str:
        """ 
        Extracts the severity level (typically a single letter or digit) from a full GHM string.

        Args:
            x (str): Full GHM string.

        Returns:
            str: Severity code.
        """
        x = x.split(" - ")[0]
        return x[-1]

    def _compute_prix_apparents(self) -> pd.DataFrame:
        """ 
        Compute average prices by GHM-GHS pair using valorisation and reimbursement data.

        Returns:
            pd.DataFrame: DataFrame with columns [GHM, GHS, prix_apparent]
        """

        data_prix_apparents = self.data_valorisations.copy()

        # Fill NaNs in GHM and preprocess to only keep 6-character code
        data_prix_apparents[self.ghm] = (
            data_prix_apparents[self.ghm]
            .ffill()
            .apply(lambda x: self._preprocess_ghm(x))
        )

        # Impute data under statistical secret
        data_prix_apparents[self.sejours_prix] = data_prix_apparents[
            self.sejours_prix
        ].replace("1 à 5", 2.5)

        data_prix_apparents[[self.sejours_prix, self.taux_am, self.montants_am]] = (
            data_prix_apparents[
                [self.sejours_prix, self.taux_am, self.montants_am]
            ].astype(float)
        )

        # Only keep data with reimbursement rates superior to 0 and stays superior to 0
        data_prix_apparents = data_prix_apparents[
            (data_prix_apparents[self.taux_am] > 0)
            & (data_prix_apparents[self.sejours_prix] > 0)
        ]

        data_prix_apparents[self.montants_am] = data_prix_apparents[
            self.montants_am
        ].fillna(0)

        # Compute the average price for a GHMxGHS
        # given by: (1/sejour) * (Montant AM / Taux AM)
        data_prix_apparents["prix_apparent"] = (
            1 / data_prix_apparents[self.sejours_prix]
        ) * (data_prix_apparents[self.montants_am] / data_prix_apparents[self.taux_am])

        return data_prix_apparents[[self.ghm, self.ghs, "prix_apparent"]]

    def _compute_prix_apparents_breakdown(
        self, breakdown: str = "racine"
    ) -> pd.DataFrame:
        """
        Computes average prices broken down a given dimension (racine, severity, type of hospitalization).

        Args:
            breakdown (str): One of `racine`, `type_hosp` or `severite`.

        Returns:
            pd.DataFrame: DataFrame with the breakdown dimension and average price.

        Raises:
            ValueError: If an unsupported breakdown type is given.
        """

        if breakdown == "racine":
            data_prix_apparents = self.data_valorisations_racine.copy()
            data_prix_apparents = data_prix_apparents[
                data_prix_apparents[self.taux_am] > 0
            ]
            data_prix_apparents["prix_apparent"] = (
                1 / data_prix_apparents[self.sejours_prix]
            ) * (
                data_prix_apparents[self.montants_am]
                / data_prix_apparents[self.taux_am]
            )
            return data_prix_apparents[[self.racine, "prix_apparent"]]

        elif breakdown == "type_hosp":
            data_prix_apparents = self.data_valorisations_type_hosp.copy()
            data_prix_apparents = data_prix_apparents[
                data_prix_apparents[self.type_hosp].isin([self.hosp_hp, self.hosp_hc])
            ]
            data_prix_apparents = data_prix_apparents[
                data_prix_apparents[self.taux_am] > 0
            ]
            data_prix_apparents["prix_apparent"] = (
                1 / data_prix_apparents[self.sejours_prix]
            ) * (
                data_prix_apparents[self.montants_am]
                / data_prix_apparents[self.taux_am]
            )
            return data_prix_apparents[[self.type_hosp, "prix_apparent"]]

        elif breakdown == "severite":
            data_prix_apparents = self.data_valorisations_severite.copy()
            data_prix_apparents = data_prix_apparents[
                data_prix_apparents[self.severite].isin(self.list_severite)
            ]
            data_prix_apparents = data_prix_apparents[
                data_prix_apparents[self.taux_am] > 0
            ]
            data_prix_apparents["prix_apparent"] = (
                1 / data_prix_apparents[self.sejours_prix]
            ) * (
                data_prix_apparents[self.montants_am]
                / data_prix_apparents[self.taux_am]
            )
            return data_prix_apparents[[self.severite, "prix_apparent"]]

        else:
            raise ValueError(f"Unsupported `breakdown` type : {breakdown}")

    def get_total_population_from_insee_estimations_by_region(self, year: int) -> float:
        """
        Retrieves total regional population from the INSEE demographic Excel file for a given year.
        INSEE data should be given by region and for 5-year age classes.

        Args:
            year (int): Year of interest.

        Returns:
            float: Total population over all regions for a given year.
        """
        
        wb = load_workbook(self.file_path_demography)
        
        ws = wb[f"{year}"]
        
        if year >= 2014:
            cell_total = "V26"
        elif year >= 1999:
            cell_total = "V25"
        elif year >= 1990:
            cell_total = "V28"
        else:
            cell_total = "V19"
                
        return ws[cell_total].value

    def get_age_class_population_from_insee_estimations_by_region(self, year: int) -> dict[str, int]:
        """
        Retrieves total population over all regions, by age classes of 15 years, to match with DIAMANT demographic data.

        Args:
            year (int) Year of interest.

        Returns:
            dict[str, int]: Dictionary with age class labels as keys and population counts as values.
        """
        wb = load_workbook(self.file_path_demography)
        
        ws = wb[f"{year}"]
        
        if year >= 2014:
            row = "26"
        elif year >= 1999:
            row = "25"
        elif year >= 1990:
            row = "28"
        else:
            row = "19"
    
        dict_pop_age_class = {}
    
        dict_pop_age_class["0-14 ans"] = sum(x for x in [ws["B"+row].value, ws["C"+row].value, ws["D"+row].value] if x is not None)
        dict_pop_age_class["15-29 ans"] = sum(x for x in [ws["E"+row].value, ws["F"+row].value, ws["G"+row].value] if x is not None)
        dict_pop_age_class["30-44 ans"] = sum(x for x in [ws["H"+row].value, ws["I"+row].value, ws["J"+row].value] if x is not None)
        dict_pop_age_class["45-59 ans"] = sum(x for x in [ws["K"+row].value, ws["L"+row].value, ws["M"+row].value] if x is not None)
        dict_pop_age_class["60-74 ans"] = sum(x for x in [ws["N"+row].value, ws["O"+row].value, ws["P"+row].value] if x is not None)
        dict_pop_age_class["75-89 ans"] = sum(x for x in [ws["Q"+row].value, ws["R"+row].value, ws["S"+row].value] if x is not None)
        dict_pop_age_class[">= 90 ans"] = sum(x for x in [ws["T"+row].value, ws["U"+row].value] if x is not None)
        
        return dict_pop_age_class
        
    def effet_volume(self) -> pd.DataFrame:
        """ 
        Computes the hospital activity using effet volume model by ATIH and its subeffects, from the casemix DIAMANT data
        and INSEE demographic data.

        Returns:
            pd.DataFrame: DataFrame containing over a period of time the total activity of French hospitals and its components.
        """

        # We compute first the effet volume and effet volume CJO
        data_volume_eco = self.data_casemix.copy()

        data_volume_eco[self.ghm] = (
            data_volume_eco[self.ghm].ffill().apply(lambda x: self._preprocess_ghm(x))
        )

        data_volume_eco[self.ghs] = data_volume_eco[self.ghs].ffill()

        data_volume_eco = data_volume_eco.drop(
            columns=[self.type_hosp, self.tranche_age]
        )

        data_volume_eco = data_volume_eco.groupby(
            by=[self.ghm, self.ghs], as_index=False
        ).sum()

        data_volume_eco = data_volume_eco.merge(
            self._compute_prix_apparents(), on=[self.ghm, self.ghs], how="right"
        )

        list_years = [col for col in data_volume_eco.columns if col.isdigit()]
        list_sorted_years = sorted(map(int, list_years))

        dict_nb_sejours = {}

        for year in list_sorted_years:
            data_volume_eco[year] = (
                data_volume_eco[str(year)] * data_volume_eco["prix_apparent"]
            )
            dict_nb_sejours[year] = data_volume_eco[str(year)].sum()

        data_volume_eco = (
            data_volume_eco[list_sorted_years].sum().to_frame("volume_economique")
        )

        data_volume_eco["effet_volume"] = data_volume_eco[
            "volume_economique"
        ].pct_change()

        data_volume_eco["effet_cjo_volume_economique"] = data_volume_eco.index.map(
            lambda x: self._effet_cjo(x)
        )

        data_volume_eco["effet_volume_cjo"] = (
            data_volume_eco["effet_volume"]
            - data_volume_eco["effet_cjo_volume_economique"]
        )

        # We can add now the effet nombre de séjours and effet structure
        data_volume_eco["nombre_sejours"] = pd.DataFrame.from_dict(
            dict_nb_sejours, orient="index", columns=["nombre_sejours"]
        )
        data_volume_eco["effet_nombre_sejours"] = data_volume_eco[
            "nombre_sejours"
        ].pct_change()

        data_volume_eco["effet_structure"] = (
            data_volume_eco["effet_volume"] - data_volume_eco["effet_nombre_sejours"]
        )

        # We also add effet nombre de séjours CJO
        data_volume_eco["effet_cjo_nombre_sejours"] = data_volume_eco.index.map(
            lambda x: self._effet_cjo(x, "sejours")
        )

        data_volume_eco["effet_nombre_sejours_cjo"] = (
            data_volume_eco["effet_nombre_sejours"]
            - data_volume_eco["effet_cjo_nombre_sejours"]
        )

        # We breakdown effet structure: effet type de prise en charge
        data_effet_prise_charge = self.data_casemix[
            [self.type_hosp] + list_years
        ].copy()
        data_effet_prise_charge[self.type_hosp] = data_effet_prise_charge[
            self.type_hosp
        ].ffill()
        data_effet_prise_charge = data_effet_prise_charge[
            data_effet_prise_charge[self.type_hosp].isin([self.hosp_hp, self.hosp_hc])
        ]
        data_effet_prise_charge = data_effet_prise_charge.groupby(
            self.type_hosp, as_index=False
        ).sum()
        data_effet_prise_charge = data_effet_prise_charge.merge(
            self._compute_prix_apparents_breakdown("type_hosp"),
            on=self.type_hosp,
            how="right",
        )

        for year in data_volume_eco.index[1:]:
            numerator = (
                data_effet_prise_charge["prix_apparent"]
                * data_effet_prise_charge[str(year)]
                / data_effet_prise_charge[str(year)].sum()
            ).sum()
            denominator = (
                data_effet_prise_charge["prix_apparent"]
                * data_effet_prise_charge[str(year - 1)]
                / data_effet_prise_charge[str(year - 1)].sum()
            ).sum()
            data_volume_eco.loc[year, "effet_type_prise_en_charge"] = (
                numerator / denominator - 1
            )

        # We breakdown effet structure : effet racine
        data_effet_racine = self.data_casemix[[self.ghm] + list_years].copy()
        data_effet_racine[self.ghm] = data_effet_racine[self.ghm].ffill()
        data_effet_racine[self.racine] = data_effet_racine[self.ghm].apply(
            lambda x: self._preprocess_racine(x)
        )
        data_effet_racine = data_effet_racine[[self.racine] + list_years]
        data_effet_racine = data_effet_racine.groupby(self.racine, as_index=False).sum()
        data_effet_racine = data_effet_racine.merge(
            self._compute_prix_apparents_breakdown("racine"),
            on=self.racine,
            how="right",
        )

        for year in data_volume_eco.index[1:]:
            numerator = (
                data_effet_racine["prix_apparent"]
                * data_effet_racine[str(year)]
                / data_effet_racine[str(year)].sum()
            ).sum()
            denominator = (
                data_effet_racine["prix_apparent"]
                * data_effet_racine[str(year - 1)]
                / data_effet_racine[str(year - 1)].sum()
            ).sum()
            data_volume_eco.loc[year, "effet_racine"] = numerator / denominator - 1

        # We breakdown effet structure : effet sévérité
        data_effet_severite = self.data_casemix[
            [self.ghm, self.type_hosp] + list_years
        ].copy()
        data_effet_severite[self.ghm] = data_effet_severite[self.ghm].ffill()
        data_effet_severite[self.type_hosp] = data_effet_severite[
            self.type_hosp
        ].ffill()
        data_effet_severite = data_effet_severite[
            data_effet_severite[self.type_hosp] == self.hosp_hc
        ]
        data_effet_severite[self.severite] = data_effet_severite[self.ghm].apply(
            lambda x: self._preprocess_severite(x)
        )
        data_effet_severite = data_effet_severite[[self.severite] + list_years]
        data_effet_severite = data_effet_severite[
            data_effet_severite[self.severite].isin(self.list_severite)
        ]
        data_effet_severite = data_effet_severite.groupby(
            self.severite, as_index=False
        ).sum()
        data_effet_severite = data_effet_severite.merge(
            self._compute_prix_apparents_breakdown("severite"),
            on=self.severite,
            how="right",
        )

        for year in data_volume_eco.index[1:]:
            numerator = (
                data_effet_severite["prix_apparent"]
                * data_effet_severite[str(year)]
                / data_effet_severite[str(year)].sum()
            ).sum()
            denominator = (
                data_effet_severite["prix_apparent"]
                * data_effet_severite[str(year - 1)]
                / data_effet_severite[str(year - 1)].sum()
            ).sum()
            data_volume_eco.loc[year, "effet_severite"] = numerator / denominator - 1

        # We breakdown effet structure : effet résiduel
        data_volume_eco["effet_residuel"] = data_volume_eco["effet_structure"] - (
            data_volume_eco["effet_type_prise_en_charge"]
            + data_volume_eco["effet_racine"]
            + data_volume_eco["effet_severite"]
        )

        # We breakdown effet nombre de séjours: effet augmentation de la population 
        data_volume_eco["effet_augmentation_population"] = (
            data_volume_eco.index.map(lambda x: 
                                     self.get_total_population_from_insee_estimations_by_region(x)
                                     /self.get_total_population_from_insee_estimations_by_region(x-1)-1)
        )
        data_volume_eco.loc[data_volume_eco.index[0], 
            "effet_augmentation_population"] = np.nan

        # We breakdown effet nombre de séjours: effet pyramide des âges
        data_pyramide_ages = self.data_casemix[[self.tranche_age]+list_years]
        data_pyramide_ages[self.tranche_age] = data_pyramide_ages[self.tranche_age].ffill()
        data_pyramide_ages = (
            data_pyramide_ages.groupby(self.tranche_age,
                                      as_index = False).sum()
        )

        for year in list_sorted_years:
            data_pyramide_ages[f"pop_{year}"] = (
                data_pyramide_ages[self.tranche_age].apply(lambda x: 
                                                          self.get_age_class_population_from_insee_estimations_by_region(year)[x])
            )

        data_volume_eco["effet_pyramide_ages"] = np.nan
        for year in data_volume_eco.index[1:]:
            numerator = (
                (data_pyramide_ages[str(year - 1)]/data_pyramide_ages[f"pop_{year-1}"])
                *(data_pyramide_ages[f"pop_{year}"]/data_pyramide_ages[f"pop_{year}"].sum())
            )
            denominator = (
                (data_pyramide_ages[str(year - 1)]/data_pyramide_ages[f"pop_{year-1}"])
                *(data_pyramide_ages[f"pop_{year-1}"]/data_pyramide_ages[f"pop_{year-1}"].sum())
            )
            data_volume_eco.loc[year, "effet_pyramide_ages"] = numerator.sum()/denominator.sum()-1
        
        # We breakdown effet nombre de séjours: effet démographie
        data_volume_eco["effet_demographie"] = (
            data_volume_eco["effet_augmentation_population"]
            + data_volume_eco["effet_pyramide_ages"]
        )
        
        # We breakdown effet nombre de séjours: effet modification du recours à l'hospitalisation
        data_volume_eco["effet_modification_recours"] = (
            data_volume_eco["effet_nombre_sejours"] - data_volume_eco["effet_demographie"]
        )

        return data_volume_eco[
            [
                "volume_economique",
                "nombre_sejours",
                "effet_volume",
                "effet_volume_cjo",
                "effet_nombre_sejours",
                "effet_nombre_sejours_cjo",
                "effet_structure",
                "effet_type_prise_en_charge",
                "effet_racine",
                "effet_severite",
                "effet_residuel",
                "effet_demographie",
                "effet_modification_recours",
                "effet_augmentation_population",
                "effet_pyramide_ages"
            ]
        ]
